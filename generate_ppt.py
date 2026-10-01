from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

output_path = r"C:\Users\kkesa\OneDrive\Desktop\New folder\Citizen_Call_Intelligence.pptx"

prs = Presentation()

# Basic color theme
NAVY = RGBColor(15, 41, 77)
BLUE = RGBColor(30, 110, 200)
SLATE = RGBColor(80, 90, 100)
LIGHT = RGBColor(242, 246, 250)


def set_title_style(slide):
    if slide.shapes.title:
        for paragraph in slide.shapes.title.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(24)
                run.font.bold = True
                run.font.color.rgb = NAVY


def add_bullet_slide(title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for idx, text in enumerate(bullets):
        p = body.paragraphs[0] if idx == 0 else body.add_paragraph()
        p.text = text
        p.level = 0
        p.bullet = True
        p.font.size = Pt(18)
        p.font.name = 'Calibri'
        p.font.color.rgb = SLATE
    set_title_style(slide)
    return slide

# Slide 1: Title
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "AI-Powered Citizen Call Intelligence Platform"
slide.placeholders[1].text = "Transforming Voice Conversations into Structured Actionable Intelligence\n\nPresenter: [Your Name / Team Name]"
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = LIGHT
for paragraph in slide.shapes.title.text_frame.paragraphs:
    for run in paragraph.runs:
        run.font.size = Pt(30)
        run.font.bold = True
        run.font.color.rgb = NAVY

# Slide 2: Challenge
add_bullet_slide(
    "The Challenge",
    [
        "Municipalities receive thousands of citizen calls every day.",
        "Calls are logged manually, causing delays, errors, and lost complaints.",
        "Departments work in silos, so recurring issues are not easily visible.",
        "Emergency cases often get delayed in routine backlogs."
    ]
)

# Slide 3: Solution
add_bullet_slide(
    "Our Solution",
    [
        "Automatically ingests live and recorded voice calls.",
        "Converts speech into transcripts and structured issue summaries.",
        "Routes complaints to the correct department in real time.",
        "Flags urgent situations for rapid action."
    ]
)

# Slide 4: Workflow
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(8.5), Inches(0.6)).text_frame.text = "Simple Workflow"

steps = [
    ("1. Call Intake", "Citizen call or audio file is received"),
    ("2. AI Processing", "Speech-to-text + categorization + sentiment"),
    ("3. Complaint Routing", "Issue assigned to correct department"),
    ("4. Action & Tracking", "Officer dashboard monitors resolution"),
]

box_w = Inches(1.8)
box_h = Inches(1.6)
start_x = Inches(0.7)
start_y = Inches(1.2)
for idx, (title, desc) in enumerate(steps):
    x = start_x + idx * (box_w + Inches(0.3))
    box = slide.shapes.add_shape(1, x, start_y, box_w, box_h)
    box.fill.solid()
    box.fill.fore_color.rgb = BLUE
    box.line.color.rgb = BLUE
    tf = box.text_frame
    tf.text = title
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = 1
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    tf.add_paragraph()
    tf.paragraphs[1].text = desc
    tf.paragraphs[1].font.size = Pt(9)
    tf.paragraphs[1].font.color.rgb = RGBColor(255, 255, 255)
    tf.paragraphs[1].alignment = 1

# Slide 5: Features
add_bullet_slide(
    "Core Features",
    [
        "Multilingual transcription for better accessibility.",
        "Department recommendation based on complaint context.",
        "Urgency and sentiment detection for priority handling.",
        "Single dashboard to track complaint lifecycle."
    ]
)

# Slide 6: Impact
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Impact"
body = slide.shapes.placeholders[1].text_frame
body.clear()
for i, text in enumerate([
    "Faster grievance resolution through automation.",
    "Improved visibility into recurring civic issues.",
    "Better accountability with clear tracking and SLA monitoring.",
    "More citizen-friendly service without requiring app-based access."
]):
    p = body.paragraphs[0] if i == 0 else body.add_paragraph()
    p.text = text
    p.level = 0
    p.bullet = True
    p.font.size = Pt(18)
    p.font.name = 'Calibri'
    p.font.color.rgb = SLATE
set_title_style(slide)

# Slide 7: Thanks
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Thank You"
body = slide.shapes.placeholders[1].text_frame
body.clear()
body.text = "AI-Powered Citizen Call Intelligence\n\nSmart Governance for Responsive Public Services\n\nQ&A"
for paragraph in body.paragraphs:
    paragraph.font.size = Pt(20)
    paragraph.font.color.rgb = SLATE
set_title_style(slide)

# Save
prs.save(output_path)
print(f"Simple presentation saved to: {output_path}")
